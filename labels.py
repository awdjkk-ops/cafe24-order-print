# -*- coding: utf-8 -*-
"""검토용 라벨 (폼텍 LS-3145: 30x9mm, 5칸 x 29줄 = 145칸)
- 라벨 엑셀: 웹 도구(cafe24-excel.html)와 같은 모양
- 바로 인쇄: 실제 mm 크기의 PDF를 축소 없이 인쇄 + 위치·배율 보정"""
import io
import re
import zipfile
from pathlib import Path

PER_PAGE = 145
COLS, ROWS = 5, 29
WEBSITE = "www.beadsmaker.com"
PINK, YELLOW, GREY, OPT_BLUE = "FFFFC0CB", "FFFFFF00", "FFBFBFBF", "FF4F81BD"

# 라벨지 실제 규격(mm) — 30x9mm 칸, 가로 간격 5mm, 세로는 틈 없음
SHEET = {"left": 20.0, "top": 18.0, "w": 30.0, "h": 9.0, "pitch_x": 35.0, "pitch_y": 9.0}
LINE_H = 3.175      # 상품명·옵션 줄 높이(mm)  (웹 도구 인쇄와 동일)
WEB_H = 2.646       # 아래 회색 띠 높이(mm)


def _t(r, k):
    v = r.get(k)
    return "" if v is None else str(v).strip()


# ---------------------------------------------------------------- 순서·색 (웹 도구와 동일)
def build_sequence(rows):
    """원본 줄(상품 단위) → 라벨 순서. 수령인(이름+주소)이 바뀔 때마다 이름 칸을 먼저."""
    seq, prev = [], None
    for r in rows:
        key = _t(r, "수령인") + "||" + _t(r, "주소")
        if key != prev:
            seq.append({"type": "NAME", "text": _t(r, "수령인"), "key": key})
            prev = key
        seq.append({"type": "PROD", "line1": _t(r, "주문상품명"), "line2": _t(r, "옵션")})
    return seq


def yellow_keys(rows):
    """같은 사람(이름+주소)이 주문을 2건 이상 → 이름 칸 노란색"""
    by = {}
    for r in rows:
        by.setdefault(_t(r, "수령인") + "||" + _t(r, "주소"), set()).add(_t(r, "주문번호"))
    return {k for k, v in by.items() if len(v) >= 2}


def name_color(item, ykeys):
    return YELLOW if item["key"] in ykeys else PINK


def paginate(seq, used_by_page=None):
    """페이지별 사용한 칸(1~145)을 건너뛰고 채움. 칸 번호는 위→아래, 왼쪽 줄부터."""
    used_by_page = used_by_page or {}
    pages, i, p = [], 0, 0
    while True:
        used = used_by_page.get(p, set())
        page = [None] * PER_PAGE
        for slot in range(1, PER_PAGE + 1):
            if slot in used:
                continue
            if i >= len(seq):
                break
            page[slot - 1] = seq[i]; i += 1
        pages.append(page)
        p += 1
        if i >= len(seq) or p > 2000:
            break
    return pages


# ---------------------------------------------------------------- 라벨 엑셀 (웹 도구와 같은 모양)
COL_WIDTHS = [1.33203125, 4.6640625, 2.44140625, 3.44140625, 0.88671875, 2.6640625, 1.33203125, 4.6640625, 2.44140625,
              3.44140625, 0.88671875, 2.109375, 1.33203125, 4.6640625, 2.44140625, 3.44140625, 0.88671875, 2.6640625,
              1.33203125, 4.6640625, 2.44140625, 3.44140625, 0.88671875, 2.33203125, 1.33203125, 4.6640625, 2.44140625,
              3.44140625, 0.88671875, 1.33203125]
ROW_HEIGHTS = [2.45, 9.0, 9.0, 7.5, 9.0, 9.0, 7.5, 9.0, 9.0, 7.5, 9.0, 9.0, 7.5, 9.0, 9.0, 7.5, 9.0, 8.25, 8.25, 9.0, 9.0,
               7.5, 9.0, 9.0, 8.25, 9.0, 9.0, 7.5, 9.0, 9.0, 7.5, 9.0, 9.0, 7.5, 9.0, 9.0, 7.5, 9.0, 9.0, 7.5, 9.0, 9.0,
               7.5, 9.0, 9.0, 7.5, 9.0, 9.75, 7.5, 9.0, 9.75, 7.5, 9.0, 9.0, 7.5, 9.0, 9.0, 7.5, 8.25, 9.0, 7.5, 9.0, 9.0,
               7.5, 9.0, 9.0, 7.5, 9.0, 9.0, 7.5, 9.0, 9.0, 7.5, 9.0, 9.0, 7.5, 9.0, 9.0, 7.5, 9.0, 9.0, 7.5, 9.0, 9.0,
               7.5, 9.0, 8.25, 7.5, 9.0, 9.0, 9.0]
ANCHORS = [0, 6, 12, 18, 24]


def write_label_xlsx(pages, ykeys, path):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.page import PageMargins
    wb = Workbook()
    wb.remove(wb.active)
    base = dict(name="돋움", size=6, family=3, charset=129)
    center_top = Alignment(horizontal="center", vertical="top", wrap_text=True)
    for p, items in enumerate(pages):
        ws = wb.create_sheet(f"--{p + 1}--")
        for i, w in enumerate(COL_WIDTHS, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        for i, h in enumerate(ROW_HEIGHTS, 1):
            ws.row_dimensions[i].height = h
        ws.page_setup.paperSize = 9
        ws.page_setup.orientation = "portrait"
        ws.page_margins = PageMargins(left=0.8267716535433072, right=0.19685039370078741, top=0.6692913385826772,
                                      bottom=0.23622047244094491, header=0.15748031496062992, footer=0.15748031496062992)
        ws.print_area = "A1:AD91"
        for col in range(COLS):
            a = 1 + ANCHORS[col]
            for g in range(ROWS):
                item = items[col * ROWS + g]
                if not item:
                    continue
                r1, r2, r3 = 2 + 3 * g, 3 + 3 * g, 4 + 3 * g
                if item["type"] == "NAME":
                    ws.merge_cells(start_row=r1, start_column=a, end_row=r2, end_column=a + 4)
                    c = ws.cell(r1, a, item["text"])
                    c.font = Font(**{**base, "size": 9}, bold=True)
                    c.alignment = center_top
                    c.fill = PatternFill("solid", fgColor=name_color(item, ykeys))
                elif not item["line2"]:
                    ws.merge_cells(start_row=r1, start_column=a, end_row=r2, end_column=a + 4)
                    c = ws.cell(r1, a, item["line1"]); c.font = Font(**base); c.alignment = center_top
                else:
                    ws.merge_cells(start_row=r1, start_column=a, end_row=r1, end_column=a + 4)
                    ws.merge_cells(start_row=r2, start_column=a, end_row=r2, end_column=a + 4)
                    c = ws.cell(r1, a, item["line1"]); c.font = Font(**base); c.alignment = center_top
                    c = ws.cell(r2, a, item["line2"]); c.font = Font(**base, bold=True, color=OPT_BLUE)
                    c.alignment = center_top
                ws.merge_cells(start_row=r3, start_column=a, end_row=r3, end_column=a + 4)
                c = ws.cell(r3, a, WEBSITE); c.font = Font(**base, bold=True); c.alignment = center_top
                c.fill = PatternFill("solid", fgColor=GREY)
    buf = io.BytesIO()
    wb.save(buf)
    # 웹 도구처럼 기본 글꼴을 '돋움 11'로 (빈 칸의 행 높이가 달라지지 않게)
    src = zipfile.ZipFile(io.BytesIO(buf.getvalue()))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        for item in src.infolist():
            data = src.read(item.filename)
            if item.filename == "xl/styles.xml":
                s = data.decode("utf-8")
                s = re.sub(r"<font>.*?</font>", '<font><sz val="11"/><name val="돋움"/><family val="3"/>'
                                                '<charset val="129"/></font>', s, count=1, flags=re.S)
                data = s.encode("utf-8")
            z.writestr(item, data)
    Path(path).write_bytes(out.getvalue())


# ---------------------------------------------------------------- 바로 인쇄용 PDF (실제 크기)
_FONTS = {}
FONT_CANDIDATES = [(r"C:\Windows\Fonts\gulim.ttc", 2), (r"C:\Windows\Fonts\malgun.ttf", None)]   # 2번 = 돋움


def _font(regular_fallback=None, bold_fallback=None):
    """돋움 등록(없으면 맑은 고딕). 반환: (보통, 굵게, 굵게를 흉내 낼지)"""
    if _FONTS:
        return _FONTS["r"], _FONTS["b"], _FONTS["fake"]
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    dotum = Path(FONT_CANDIDATES[0][0])
    if dotum.exists():
        try:
            pdfmetrics.registerFont(TTFont("LBL", str(dotum), subfontIndex=2))     # 돋움
            _FONTS.update(r="LBL", b="LBL", fake=True)          # 돋움은 굵은 파일이 없어 웹처럼 굵게 흉내
            return _FONTS["r"], _FONTS["b"], _FONTS["fake"]
        except Exception:
            pass
    reg = next((p for p in (r"C:\Windows\Fonts\malgun.ttf", regular_fallback) if p and Path(p).exists()), None)
    bold = next((p for p in (r"C:\Windows\Fonts\malgunbd.ttf", bold_fallback) if p and Path(p).exists()), None)
    if not reg:
        raise RuntimeError("라벨용 글꼴을 찾지 못했습니다.")
    pdfmetrics.registerFont(TTFont("LBL", reg))
    if bold:
        pdfmetrics.registerFont(TTFont("LBLB", bold))
    _FONTS.update(r="LBL", b="LBLB" if bold else "LBL", fake=not bold)
    return _FONTS["r"], _FONTS["b"], _FONTS["fake"]


def _hex(argb):
    from reportlab.lib.colors import HexColor
    return HexColor("#" + argb[2:])


def _fit(c, text, font, size, width):
    """한 줄에 들어가는 만큼 자르기"""
    from reportlab.pdfbase.pdfmetrics import stringWidth
    if stringWidth(text, font, size) <= width:
        return text
    while text and stringWidth(text + "…", font, size) > width:
        text = text[:-1]
    return text


def _wrap2(text, font, size, width):
    """글자 단위로 최대 2줄 (웹 도구의 word-break:break-all과 같음)"""
    from reportlab.pdfbase.pdfmetrics import stringWidth
    lines, cur = [], ""
    for ch in text:
        if stringWidth(cur + ch, font, size) > width:
            lines.append(cur); cur = ch
            if len(lines) == 2:
                break
        else:
            cur += ch
    if len(lines) < 2 and cur:
        lines.append(cur)
    return lines[:2]


def _geom(cal):
    """보정값 반영한 칸 위치 계산 함수"""
    ox, oy = float(cal.get("offset_x", 0)), float(cal.get("offset_y", 0))
    sx, sy = 1 + float(cal.get("scale_x", 0)) / 100, 1 + float(cal.get("scale_y", 0)) / 100

    def slot(col, g):
        x = (SHEET["left"] + col * SHEET["pitch_x"]) * sx + ox
        y = (SHEET["top"] + g * SHEET["pitch_y"]) * sy + oy
        return x, y, SHEET["w"] * sx, SHEET["h"] * sy
    return slot


def write_label_pdf(pages, ykeys, path, cal=None, fonts=(None, None)):
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import mm
    from reportlab.pdfgen import canvas
    cal = cal or {}
    fr, fb, fake = _font(*fonts)
    slot = _geom(cal)
    W, H = A4
    c = canvas.Canvas(str(path), pagesize=A4)
    c.setTitle("검토용 라벨")

    def text(fn, x_, y_, s_, bold=False):
        if not (bold and fake):
            fn(x_, y_, s_)
            return
        from reportlab.pdfbase.pdfmetrics import stringWidth
        if fn == c.drawCentredString:
            x_ -= stringWidth(s_, c._fontname, c._fontsize) / 2
        t = c.beginText()
        t.setTextRenderMode(2)                 # 채우기+테두리 → 굵게 보임
        t.setTextOrigin(x_, y_)
        t.setFont(c._fontname, c._fontsize)
        c.saveState()
        c.setStrokeColor(c._fillColorObj); c.setLineWidth(0.22)
        t.textOut(s_)
        c.drawText(t)
        c.restoreState()
    for items in pages:
        for col in range(COLS):
            for g in range(ROWS):
                item = items[col * ROWS + g]
                if not item:
                    continue
                x, y, w, h = slot(col, g)
                top = H - y * mm                                   # PDF는 아래가 0
                k = h / SHEET["h"]
                web_h, text_h = WEB_H * k, (LINE_H * 2) * k
                # 아래 회색 띠 + 사이트 주소
                c.setFillColor(_hex(GREY))
                c.rect(x * mm, top - h * mm, w * mm, web_h * mm, stroke=0, fill=1)
                c.setFillColor(_hex("FF000000")); c.setFont(fb, 6)
                text(c.drawCentredString, (x + w / 2) * mm, top - h * mm + web_h * mm * 0.28, WEBSITE, bold=True)
                inner = (w - 1.0) * mm
                if item["type"] == "NAME":
                    c.setFillColor(_hex(name_color(item, ykeys)))
                    c.rect(x * mm, top - text_h * mm, w * mm, text_h * mm, stroke=0, fill=1)
                    c.setFillColor(_hex("FF000000")); c.setFont(fb, 8)
                    text(c.drawCentredString, (x + w / 2) * mm, top - 0.4 * mm - 8 * 0.9,
                         _fit(c, item["text"], fb, 8, inner), bold=True)
                elif not item["line2"]:
                    c.setFont(fr, 6)
                    for n, line in enumerate(_wrap2(item["line1"], fr, 6, inner)):
                        c.drawCentredString((x + w / 2) * mm, top - 0.4 * mm - 6 * 0.9 - n * 6 * 1.25, line)
                else:
                    c.setFont(fr, 6)
                    c.drawString((x + 1.0) * mm, top - 0.4 * mm - 6 * 0.9, _fit(c, item["line1"], fr, 6, inner - mm))
                    c.setFillColor(_hex(OPT_BLUE)); c.setFont(fb, 6)
                    text(c.drawCentredString, (x + w / 2) * mm, top - LINE_H * k * mm - 0.3 * mm - 6 * 0.9,
                         _fit(c, item["line2"], fb, 6, inner), bold=True)
                    c.setFillColor(_hex("FF000000"))
        c.showPage()
    c.save()


def write_calibration_pdf(path, cal=None, fonts=(None, None)):
    """맞춤 확인용: 145칸 테두리 + 칸 번호. 일반 A4에 인쇄해 라벨지와 겹쳐 불빛에 비춰보세요."""
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import mm
    from reportlab.pdfgen import canvas
    cal = cal or {}
    fr, fb, _fake = _font(*fonts)
    slot = _geom(cal)
    W, H = A4
    c = canvas.Canvas(str(path), pagesize=A4)
    c.setLineWidth(0.25)
    for col in range(COLS):
        for g in range(ROWS):
            x, y, w, h = slot(col, g)
            top = H - y * mm
            c.rect(x * mm, top - h * mm, w * mm, h * mm, stroke=1, fill=0)
            c.setFont(fr, 5)
            c.drawCentredString((x + w / 2) * mm, top - h * mm * 0.62, str(col * ROWS + g + 1))
    c.setFont(fb, 7)
    c.drawString(20 * mm, H - 8 * mm,
                 f"라벨 맞춤 확인 · 위치 보정 가로 {cal.get('offset_x', 0)}mm 세로 {cal.get('offset_y', 0)}mm · "
                 f"배율 보정 가로 {cal.get('scale_x', 0)}% 세로 {cal.get('scale_y', 0)}%")
    c.setFont(fr, 6)
    c.drawString(20 * mm, H - 12 * mm, "이 종이를 라벨지와 겹쳐 불빛에 비춰보고, 칸 테두리가 라벨 칸과 맞는지 확인하세요. "
                                       "(인쇄 설정: 실제 크기 100%)")
    c.showPage()
    c.save()
