# -*- coding: utf-8 -*-
"""세금계산서 발행 — 데이터 (사업자 목록, 요청, 발행 내역, 홈택스 일괄발급 엑셀, 기존 엑셀 가져오기)

모든 내용은 이 PC의 data/tax_invoice.json 한 파일에 저장합니다 (업데이트로 지워지지 않는 폴더).
하루에 한 번 data/tax_backup/ 에 이전 내용을 복사해 둡니다 (최근 60개)."""
import datetime as dt
import re
import shutil
import uuid
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

import order_print as op

TAX_FILE = op.DATA / "tax_invoice.json"
BACKUP_DIR = op.DATA / "tax_backup"
OUT_DIR = op.BASE_DIR / "세금계산서 파일"
BACKUP_KEEP = 60

ITEM_NAME = "악세사리 부자재"
ITEM_SPEC = "세트1"

# 우리 회사(공급자) 기본값 — 탭의 [우리 회사 정보]에서 고칠 수 있음
SUPPLIER_DEFAULT = {"bizno": "1282895100", "name": "엔비즈", "ceo": "정지훈",
                    "addr": "경기도 고양시 일산동구 중앙로1261번길 77, 1002호(장항동, 메트로폴리스)",
                    "uptae": "소매업", "jongmok": "통신판매업", "email": "reims@naver.com"}

BIZ_FIELDS = (("bizno", "사업자번호"), ("name", "상호"), ("ceo", "성명"), ("addr", "주소"),
              ("uptae", "업태"), ("jongmok", "종목"), ("email", "이메일"))


# ---------------------------------------------------------------- 저장
def _empty():
    return {"version": 1, "supplier": dict(SUPPLIER_DEFAULT), "biz": {}, "lines": [], "invoices": [], "groups": {}}


def load():
    d = op.read_json(TAX_FILE, None) or _empty()
    for k, v in _empty().items():
        d.setdefault(k, v)
    return d


def save(d):
    op.DATA.mkdir(parents=True, exist_ok=True)
    _daily_backup()
    op.write_json(TAX_FILE, d)


def _daily_backup():
    """오늘 처음 저장할 때 지금까지의 내용을 백업 폴더에 복사"""
    if not TAX_FILE.exists():
        return
    BACKUP_DIR.mkdir(parents=True, exist_ok=True)
    dst = BACKUP_DIR / f"tax_invoice_{dt.date.today():%Y%m%d}.json"
    if not dst.exists():
        shutil.copy2(TAX_FILE, dst)
        olds = sorted(BACKUP_DIR.glob("tax_invoice_*.json"))
        for p in olds[:-BACKUP_KEEP]:
            p.unlink(missing_ok=True)


def new_id():
    return uuid.uuid4().hex[:12]


def now():
    return dt.datetime.now().isoformat(timespec="seconds")


# ---------------------------------------------------------------- 사업자번호·금액
def digits(v):
    return re.sub(r"\D", "", str(v or ""))


def bizno_ok(v):
    """사업자등록번호 검사 (국세청 검증식)"""
    d = digits(v)
    if len(d) != 10:
        return False
    n = [int(c) for c in d]
    s = sum(a * b for a, b in zip(n[:9], (1, 3, 7, 1, 3, 7, 1, 3, 5)))
    s += (n[8] * 5) // 10
    return (10 - s % 10) % 10 == n[9]


def bizno_fmt(v):
    d = digits(v)
    return f"{d[:3]}-{d[3:5]}-{d[5:]}" if len(d) == 10 else str(v or "")


def split_vat(total):
    """결제금액(부가세 포함) → (공급가액, 세액). 홈택스 [계산] 버튼과 같은 방식 (공급가액 반올림)"""
    total = int(total)
    supply = int((Decimal(total) / Decimal("1.1")).quantize(Decimal("1"), rounding=ROUND_HALF_UP))
    return supply, total - supply


def won(v):
    try:
        return f"{int(v):,}"
    except (TypeError, ValueError):
        return str(v or "")


def to_int(v):
    s = str(v or "").replace(",", "").replace("원", "").strip()
    if not s:
        return 0
    return int(round(float(s)))


def parse_date(v):
    """'2026-10-13', '20261013', '2026.10.13', datetime → date (못 읽으면 None)"""
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    d = digits(str(v)[:10]) if v else ""
    if len(d) >= 8:
        try:
            return dt.date(int(d[:4]), int(d[4:6]), int(d[6:8]))
        except ValueError:
            return None
    return None


def oid_date(oid):
    """카페24 주문번호 앞 8자리(주문일)"""
    s = str(oid or "")
    return parse_date(s[:8]) if s[:8].isdigit() else None


# ---------------------------------------------------------------- 사업자
def biz_missing(b):
    """비어 있는 항목 이름 목록"""
    if not b:
        return [n for _, n in BIZ_FIELDS]
    return [n for k, n in BIZ_FIELDS if not str(b.get(k) or "").strip()]


def search_biz(d, text):
    """고객명·상호·성명·사업자번호·이메일 일부로 찾기"""
    t = text.strip().lower()
    td = digits(t)
    res = []
    for b in d["biz"].values():
        hay = " ".join([b.get("name", ""), b.get("ceo", ""), b.get("email", ""), " ".join(b.get("customers", []))]).lower()
        if not t or t in hay or (len(td) >= 3 and td in digits(b.get("bizno"))):
            res.append(b)
    res.sort(key=lambda b: b.get("updated", ""), reverse=True)
    return res


def find_biz_by_no(d, bizno):
    dn = digits(bizno)
    return next((b for b in d["biz"].values() if digits(b.get("bizno")) == dn), None) if dn else None


def save_biz(d, info, biz_id=None, customer=None):
    """새로 만들거나 고침. 같은 사업자번호가 이미 있으면 그쪽을 고침. 돌려줌: 사업자 id"""
    customers = info.get("customers")
    memo = info.get("memo")
    info = {k: str(info.get(k) or "").strip() for k, _ in BIZ_FIELDS}
    info["bizno"] = digits(info["bizno"])
    b = d["biz"].get(biz_id) if biz_id else None
    if not b:
        b = find_biz_by_no(d, info["bizno"])
    if not b:
        b = {"id": new_id(), "customers": [], "memo": "", "created": now()}
        d["biz"][b["id"]] = b
    b.update(info)
    if customers is not None:
        b["customers"] = [c for c in customers if c]
    if memo is not None:
        b["memo"] = str(memo).strip()
    if customer and customer not in b["customers"]:
        b["customers"].append(customer)
    b["updated"] = now()
    return b["id"]


def biz_usage(d, biz_id):
    return sum(1 for l in d["lines"] if l.get("biz_id") == biz_id)


# ---------------------------------------------------------------- 요청(주문 줄)과 계산서 묶음
def add_line(d, **kw):
    line = {"id": new_id(), "status": "wait", "customer": "", "company": "", "biz_id": "", "oid": "",
            "date": "", "amount": 0, "memo": "", "pay": "", "split": False, "excel_at": "", "created": now()}
    line.update(kw)
    d["lines"].append(line)
    return line


def line_by_id(d, lid):
    return next((l for l in d["lines"] if l["id"] == lid), None)


def month_of(line):
    return str(line.get("date") or "")[:7]


def group_key(line):
    who = line.get("biz_id") or ("name:" + (line.get("company") or line.get("customer") or "?"))
    key = f"{who}|{month_of(line)}"
    return key + f"|{line['id']}" if line.get("split") else key


def waiting_groups(d):
    """대기 중인 줄을 계산서 한 장 단위로 묶음. 오래된 달 먼저"""
    groups = {}
    for l in d["lines"]:
        if l["status"] == "wait":
            groups.setdefault(group_key(l), []).append(l)
    out = [invoice_view(d, k, ls) for k, ls in groups.items()]
    out.sort(key=lambda g: (g["month"], g["company"], g["customer"]))
    return out


def invoice_view(d, key, lines):
    """계산서 한 장에 들어갈 값과 문제점"""
    lines = sorted(lines, key=lambda l: (l.get("date") or "", l.get("oid") or ""))
    opt = d["groups"].get(key, {})
    dates = [parse_date(l.get("date")) for l in lines if parse_date(l.get("date"))]
    mode = opt.get("mode", "last")
    if mode == "custom" and parse_date(opt.get("date")):
        wdate = parse_date(opt["date"])
    elif dates:
        wdate = min(dates) if mode == "first" else max(dates)
    else:
        wdate = None
    total = sum(int(l.get("amount") or 0) for l in lines)
    supply, tax = split_vat(total)
    biz = d["biz"].get(lines[0].get("biz_id") or "")
    problems = []
    if not biz:
        problems.append("사업자 정보 없음")
    else:
        miss = biz_missing(biz)
        if miss:
            problems.append("빈 칸: " + ", ".join(miss))
        elif not bizno_ok(biz["bizno"]):
            problems.append("사업자번호 확인 필요")
    if not wdate:
        problems.append("작성일자 없음")
    elif dates and any(x.strftime("%Y-%m") != wdate.strftime("%Y-%m") for x in dates):
        problems.append("작성일자의 달이 주문 달과 다름")
    if total <= 0:
        problems.append("금액이 0원 이하")
    cards = [l for l in lines if "카드" in str(l.get("pay") or "")]
    return {"key": key, "lines": lines, "biz": biz, "date": wdate, "mode": mode, "total": total, "supply": supply,
            "tax": tax, "problems": problems, "cards": len(cards), "month": month_of(lines[0]),
            "customer": lines[0].get("customer") or (biz or {}).get("ceo", ""),
            "company": (biz or {}).get("name") or lines[0].get("company") or "",
            "excel_at": max((l.get("excel_at") or "" for l in lines), default="")}


def set_group_date(d, key, mode, date=None):
    if mode in ("last", ""):
        d["groups"].pop(key, None)
    else:
        d["groups"][key] = {"mode": mode, "date": date.isoformat() if date else ""}


def mark_done(d, group, how):
    """대기 묶음 → 발행 완료 (그때의 사업자 정보를 함께 보관)"""
    inv = {"id": new_id(), "issued_at": now(), "how": how,
           "date": group["date"].isoformat() if group["date"] else "",
           "total": group["total"], "supply": group["supply"], "tax": group["tax"],
           "biz": dict(group["biz"] or {}), "customer": group["customer"], "company": group["company"],
           "line_ids": [l["id"] for l in group["lines"]], "imported": False}
    d["invoices"].append(inv)
    for l in group["lines"]:
        l["status"], l["inv"] = "done", inv["id"]
    d["groups"].pop(group["key"], None)
    return inv


def undo_done(d, inv_id):
    """발행 완료를 취소하고 대기로 되돌림"""
    inv = next((i for i in d["invoices"] if i["id"] == inv_id), None)
    if not inv:
        return 0
    n = 0
    for l in d["lines"]:
        if l.get("inv") == inv_id:
            l["status"] = "wait"; l.pop("inv", None); n += 1
    d["invoices"].remove(inv)
    return n


def done_invoices(d):
    by = {}
    for l in d["lines"]:
        if l.get("inv"):
            by.setdefault(l["inv"], []).append(l)
    out = []
    for inv in d["invoices"]:
        out.append(dict(inv, lines=sorted(by.get(inv["id"], []), key=lambda l: (l.get("date") or "", l.get("oid") or ""))))
    out.sort(key=lambda i: (i.get("date") or "", i.get("issued_at") or ""), reverse=True)
    return out


def deadline_info(d, today=None):
    """지난달까지의 대기 건 → (건수, 마감일, 지났는지)"""
    today = today or dt.date.today()
    this_month = today.strftime("%Y-%m")
    old = [g for g in waiting_groups(d) if g["month"] and g["month"] < this_month]
    due = dt.date(today.year, today.month, 10)
    return len(old), due, today > due


# ---------------------------------------------------------------- 카페24 주문 → 줄
def order_info(o):
    buyer = o.get("buyer") or {}
    if isinstance(buyer, list):
        buyer = buyer[0] if buyer else {}
    amt = o.get("actual_order_amount") or {}
    pay = o.get("payment_amount")
    if pay in (None, ""):
        pay = amt.get("payment_amount") if isinstance(amt, dict) else 0
    ship = amt.get("shipping_fee") if isinstance(amt, dict) else 0
    names = o.get("payment_method_name") or o.get("payment_method") or []
    names = names if isinstance(names, list) else [names]
    methods = o.get("payment_method") or []
    methods = methods if isinstance(methods, list) else [methods]
    pay_txt = ", ".join(str(n) for n in names if n)
    if "card" in methods and "카드" not in pay_txt:
        pay_txt = (pay_txt + ", 신용카드").strip(", ")
    return {"oid": o.get("order_id", ""), "date": str(o.get("order_date") or "")[:10],
            "customer": buyer.get("name") or o.get("buyer_name") or "",
            "amount": to_int(pay), "shipping": to_int(ship), "pay": pay_txt,
            "canceled": o.get("canceled") == "T"}


def fetch_order(cfg, oid):
    token = op.get_access_token(cfg)
    o = op._fetch_one(cfg, token, oid.strip())
    return order_info(o) if o else None


def lines_with_oid(d, oid):
    return [l for l in d["lines"] if oid and l.get("oid") == oid]


# ---------------------------------------------------------------- 홈택스 일괄발급 엑셀
HT_TITLE = "엑셀 업로드 양식(전자세금계산서-일반(영세율)) - 100건 이하"
HT_NOTES = [
    "○ 필수항목(주황색)은 반드시 입력하셔야 합니다.",
    "○ 임의로 양식을 변경[행 또는 열 추가 삭제 등]하는 경우 발급시 오류가 발생할 수 있으므로, 정해진 양식으로 작성하시기 바랍니다\n"
    "     > 실제 업로드할 DATA는 7행부터 입력하여야 하며, 최대 100건까지 입력이 가능합니다.(100건 초과 자료는 처리 안되며, 발급은 최대 50건씩 처리가능합니다)",
    "○ 이 파일은 '주문서 출력 관리' 프로그램이 홈택스 양식에 맞춰 만들었습니다.",
    "○ 발급가능한 파일 확장자는 XLS, XLSX 입니다.",
]
HT_HEADERS = (["전자(세금)계산서 종류\n(01:일반, 02:영세율)", "작성일자", "공급자 등록번호\n(\"-\" 없이 입력)",
               "공급자\n 종사업장번호", "공급자 상호", "공급자 성명", "공급자 사업장주소", "공급자 업태", "공급자 종목",
               "공급자 이메일", "공급받는자 등록번호\n(\"-\" 없이 입력)", "공급받는자 \n종사업장번호", "공급받는자 상호 ",
               "공급받는자 성명", "공급받는자 사업장주소", "공급받는자 업태", "공급받는자 종목", "공급받는자 이메일1",
               "공급받는자 이메일2", "공급가액\n합계", "세액\n합계", "비고"]
              + [h for i in range(1, 5) for h in (f"일자{i}\n(2자리, 작성년월 제외)", f"품목{i}", f"규격{i}", f"수량{i}",
                                                   f"단가{i}", f"공급가액{i}", f"세액{i}", f"품목비고{i}")]
              + ["현금", "수표", "어음", "외상미수금", "영수(01),\n청구(02)"])
HT_PER_FILE = 100


def hometax_row(sup, g):
    b, day = g["biz"], g["date"]
    item = [f"{day.day:02d}", ITEM_NAME, ITEM_SPEC, "1", str(g["supply"]), str(g["supply"]), str(g["tax"]), ""]
    row = (["01", day.strftime("%Y%m%d"), digits(sup["bizno"]), "", sup["name"], sup["ceo"], sup["addr"],
            sup["uptae"], sup["jongmok"], sup["email"],
            digits(b["bizno"]), "", b["name"], b["ceo"], b["addr"], b["uptae"], b["jongmok"], b["email"], "",
            str(g["supply"]), str(g["tax"]), ""]
           + item + [""] * 24 + ["", "", "", "", "01"])
    assert len(row) == len(HT_HEADERS) == 59
    return row


def make_hometax_excel(d, groups, folder=None):
    """묶음 목록 → 홈택스 일괄발급용 xlsx (100장씩 나눔). 돌려줌: 파일 경로 목록"""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    folder = Path(folder or OUT_DIR)
    folder.mkdir(parents=True, exist_ok=True)
    stamp = dt.datetime.now().strftime("%Y%m%d_%H%M")
    paths = []
    chunks = [groups[i:i + HT_PER_FILE] for i in range(0, len(groups), HT_PER_FILE)]
    for n, chunk in enumerate(chunks, 1):
        wb = Workbook()
        ws = wb.active
        ws.title = "엑셀업로드양식"
        ws.cell(1, 1, HT_TITLE).font = Font(bold=True, size=12)
        for i, t in enumerate(HT_NOTES, 2):
            ws.cell(i, 1, t).alignment = Alignment(wrap_text=False, vertical="top")
        orange = PatternFill("solid", fgColor="FFC000")
        required = {0, 1, 2, 4, 5, 10, 12, 13, 19, 20, 58}
        for c, h in enumerate(HT_HEADERS, 1):
            cell = ws.cell(6, c, h)
            cell.font = Font(bold=True, size=9)
            cell.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")
            if c - 1 in required:
                cell.fill = orange
            ws.column_dimensions[get_column_letter(c)].width = 14
        ws.row_dimensions[6].height = 42
        for r, g in enumerate(chunk, 7):
            for c, v in enumerate(hometax_row(d["supplier"], g), 1):
                cell = ws.cell(r, c, v)
                cell.number_format = "@"          # 모두 글자로 (앞자리 0 유지)
        suffix = f"_{n}" if len(chunks) > 1 else ""
        p = folder / f"홈택스_일괄발급_{stamp}{suffix}.xlsx"
        wb.save(p)
        paths.append(p)
    return paths


def mark_excel_made(d, groups):
    t = now()
    for g in groups:
        for l in g["lines"]:
            l["excel_at"] = t


# ---------------------------------------------------------------- 전체 내보내기 (백업용 엑셀)
def export_all(d, path=None):
    from openpyxl import Workbook
    from openpyxl.styles import Font
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    path = Path(path or OUT_DIR / f"세금계산서_전체내역_{dt.datetime.now():%Y%m%d_%H%M}.xlsx")
    wb = Workbook()
    ws = wb.active
    ws.title = "사업자 목록"
    ws.append(["사업자번호", "상호", "성명", "주소", "업태", "종목", "이메일", "고객명", "메모"])
    for b in sorted(d["biz"].values(), key=lambda b: b.get("name", "")):
        ws.append([bizno_fmt(b.get("bizno")), b.get("name"), b.get("ceo"), b.get("addr"), b.get("uptae"),
                   b.get("jongmok"), b.get("email"), ", ".join(b.get("customers", [])), b.get("memo", "")])
    ws2 = wb.create_sheet("발행 내역")
    ws2.append(["상태", "작성일자", "고객명", "상호", "사업자번호", "주문번호", "주문일", "금액", "메모",
                "총발행금액", "공급가액", "세액", "발행 방법", "완료 처리한 때"])
    for inv in done_invoices(d):
        b = inv.get("biz") or {}
        for i, l in enumerate(inv["lines"] or [{}]):
            ws2.append(["발행완료", inv.get("date"), inv.get("customer"), inv.get("company"), bizno_fmt(b.get("bizno")),
                        l.get("oid"), l.get("date"), l.get("amount"), l.get("memo"),
                        inv.get("total") if i == 0 else None, inv.get("supply") if i == 0 else None,
                        inv.get("tax") if i == 0 else None, inv.get("how"), str(inv.get("issued_at", ""))[:16]])
    for g in waiting_groups(d):
        for i, l in enumerate(g["lines"]):
            ws2.append(["대기", g["date"].isoformat() if g["date"] else "", g["customer"], g["company"],
                        bizno_fmt((g["biz"] or {}).get("bizno")), l.get("oid"), l.get("date"), l.get("amount"),
                        l.get("memo"), g["total"] if i == 0 else None, g["supply"] if i == 0 else None,
                        g["tax"] if i == 0 else None, "", ""])
    for w in (ws, ws2):
        for c in w[1]:
            c.font = Font(bold=True)
    wb.save(path)
    return path


# ---------------------------------------------------------------- 기존 엑셀 가져오기
def _sheet_layout(ws):
    """시트 머리줄을 보고 칸 위치를 정함 (0부터)"""
    head = [str(c.value or "").strip() for c in ws[1]]
    if head and head[0] == "작성일자":          # '2025년 이전' 시트
        return {"wdate": 0, "customer": 1, "company": 2, "oid": 3, "amount": 4, "memo": 5, "total": 6,
                "email": 7, "bizno": 8, "name": 9, "ceo": 10, "addr": 11, "uptae": 12, "jongmok": 13}
    if head and head[0] == "고객명":             # '2026년' 시트
        return {"wdate": None, "customer": 0, "company": 1, "oid": 2, "amount": 3, "memo": 4, "total": 5,
                "bizno": 6, "name": 7, "ceo": 8, "addr": 9, "uptae": 10, "jongmok": 11, "email": 12}
    return None


def _read_groups(ws, lay):
    """병합된 칸(고객명·발행금액·작성일자)을 기준으로 계산서 한 장씩 묶어서 읽기"""
    key_cols = {lay[k] + 1 for k in ("customer", "total", "wdate") if lay.get(k) is not None}
    span = {}
    for rng in ws.merged_cells.ranges:
        if rng.min_col in key_cols and rng.max_col == rng.min_col:
            for r in range(rng.min_row, rng.max_row + 1):
                span.setdefault(r, rng.min_row)
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    groups, cur = [], None
    for idx, r in enumerate(rows, start=2):
        has_order = r[lay["oid"]] not in (None, "") or r[lay["amount"]] not in (None, "")
        if not has_order:
            continue
        start = span.get(idx, idx)
        if cur is None or start != cur["start"]:
            # 병합 밖인데 고객명·발행금액·작성일자가 모두 비어 있으면 앞 묶음에 이어 붙임
            top = r if start == idx else rows[start - 2]
            blank = all(top[lay[k]] in (None, "") for k in ("customer", "total") ) and \
                (lay["wdate"] is None or top[lay["wdate"]] in (None, ""))
            if cur is not None and start == idx and blank and idx not in span:
                cur["rows"].append(r)
                continue
            cur = {"start": start, "top": top, "rows": []}
            groups.append(cur)
        cur["rows"].append(r)
    return groups


def _cell(r, i):
    if i is None or i >= len(r):
        return ""
    v = r[i]
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def import_excel(d, path, replace_old=True):
    """기존 세금계산서 엑셀(2026년 / 2025년 이전 시트)을 가져옴.
    - 사업자 목록: 7칸이 모두 있는 사업자만 (같은 번호가 여러 번이면 나중 것)
    - 발행 내역: 모든 묶음 (사업자 정보가 없으면 '사업자 정보 없음'으로)"""
    from openpyxl import load_workbook
    wb = load_workbook(path, data_only=False)
    sheets = []
    for ws in wb.worksheets:
        lay = _sheet_layout(ws)
        if lay:
            sheets.append((ws, lay))
    if not sheets:
        raise ValueError("가져올 시트를 찾지 못했습니다 (첫 줄이 '작성일자' 또는 '고객명'인 시트)")
    # 오래된 시트부터 (나중 정보가 앞의 정보를 덮도록)
    sheets.sort(key=lambda s: 0 if s[1]["wdate"] is not None else 1)
    if replace_old:
        old_inv = {i["id"] for i in d["invoices"] if i.get("imported")}
        d["invoices"] = [i for i in d["invoices"] if not i.get("imported")]
        d["lines"] = [l for l in d["lines"] if l.get("inv") not in old_inv]
    stat = {"groups": 0, "lines": 0, "biz_new": 0, "biz_upd": 0, "partial": 0, "no_biz": 0, "linked": 0}
    pending = []
    for ws, lay in sheets:
        for g in _read_groups(ws, lay):
            top = g["top"]
            # 사업자 정보는 묶음의 첫 줄이 아닌 줄에 적혀 있기도 함 → 칸마다 처음 나오는 값
            info = {k: next((_cell(r, lay[k]) for r in [top] + g["rows"] if _cell(r, lay[k])), "")
                    for k, _ in BIZ_FIELDS}
            info["bizno"] = digits(info["bizno"])
            filled = [k for k, _ in BIZ_FIELDS if info[k]]
            complete = len(filled) == len(BIZ_FIELDS) and "@" in info["email"] and len(info["bizno"]) == 10
            customer, company = _cell(top, lay["customer"]), _cell(top, lay["company"])
            biz_id = ""
            if complete:
                existed = find_biz_by_no(d, info["bizno"])
                biz_id = save_biz(d, info, customer=customer or None)
                stat["biz_upd" if existed else "biz_new"] += 1
            elif any(info[k] for k in ("bizno", "name", "ceo", "addr", "uptae", "jongmok")):
                stat["partial"] += 1
            lines = []
            for r in g["rows"]:
                oid = _cell(r, lay["oid"])
                try:
                    amount = to_int(r[lay["amount"]])
                except ValueError:
                    amount = 0
                od = oid_date(oid)
                lines.append({"oid": oid, "amount": amount, "memo": _cell(r, lay["memo"]),
                              "date": od.isoformat() if od else ""})
            wdate = parse_date(top[lay["wdate"]]) if lay["wdate"] is not None else None
            if not wdate:
                ds = [parse_date(l["date"]) for l in lines if l["date"]]
                wdate = max(ds) if ds else None
            pending.append((biz_id, info if complete else None, customer, company, wdate, lines, ws.title))
    # 사업자 정보 없는 묶음은 상호·고객명이 같은 등록 사업자에 연결
    by_name = {}
    for b in d["biz"].values():
        for nm in [b.get("name")] + b.get("customers", []):
            if nm:
                by_name.setdefault(nm.replace(" ", ""), b)
    for biz_id, info, customer, company, wdate, lines, sheet in pending:
        biz = d["biz"].get(biz_id) if biz_id else None
        if not biz:
            biz = by_name.get((company or "").replace(" ", "")) or by_name.get((customer or "").replace(" ", ""))
            if biz:
                stat["linked"] += 1
            else:
                stat["no_biz"] += 1
        total = sum(l["amount"] for l in lines)
        supply, tax = split_vat(total)
        inv = {"id": new_id(), "issued_at": "", "how": f"기존 엑셀 ({sheet})", "date": wdate.isoformat() if wdate else "",
               "total": total, "supply": supply, "tax": tax, "biz": dict(info or biz or {}) if (info or biz) else {},
               "customer": customer, "company": company or (biz or {}).get("name", ""), "line_ids": [], "imported": True}
        for l in lines:
            line = add_line(d, status="done", customer=customer, company=company, biz_id=(biz or {}).get("id", ""),
                            oid=l["oid"], date=l["date"], amount=l["amount"], memo=l["memo"], inv=inv["id"],
                            created=now())
            inv["line_ids"].append(line["id"])
        d["invoices"].append(inv)
        stat["groups"] += 1
        stat["lines"] += len(lines)
    stat["biz_total"] = len(d["biz"])
    return stat
