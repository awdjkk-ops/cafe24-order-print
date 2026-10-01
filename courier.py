# -*- coding: utf-8 -*-
"""택배 발송용 파일(롯데택배 알프스) + 확인 목록 — 웹 도구(cafe24-excel.html)와 같은 규칙으로 만듦.
입력은 카페24 '엑셀 바로 다운로드'와 같은 형식의 줄(열 이름 → 값 dict) 목록."""
import json
import re
import time

import requests

# ---------------------------------------------------------------- 고정값 (웹 도구와 동일)
COURIER_HEADERS = [
    "받는분성명", "받는분우편번호", "받는분주소(전체, 분할)", "받는분전화번호1(P)", "받는분전화번호2(Q)",
    "상품명(옵션명)", "수량", "자동합계\n총 수량",
    "사용안함\n주문번호 판매처등활용시\n씨엔플러스에만노출",
    "배송메세지1®", "운임구분", "기본운임", "보내는분성명", "보내는분전화번호1", "보내는분전화번호2", "보내는분주소(전체,분할)",
]
COURIER_WIDTHS = [22, 10, 42, 16, 16, 20, 7, 10, 14, 20, 10, 10, 12, 16, 16, 42]
SENDER = {"name": "엔비즈", "tel1": "031-915-5123", "tel2": "031-915-5123",
          "addr": "서울 서초구 양재동 230, 롯데택배 양재대리점1 엔비즈"}
YELLOW, RED, BLUE = "FFFFFF00", "FFC0392B", "FF1F4E9C"

TRIVIAL_DELIVERY_MESSAGES = {
    "문 앞", "문앞", "직접 받고 부재 시 문 앞", "0", "현관", "경비실",
    "문 앞에 놓아주세요", "빠른 발송 부탁드립니다", "문앞에 놔주세요", "문앞에 놓아주십시요",
    "배송전 문자주세요", "배송 전 미리 연락해 주세요", "문 앞에 놓아주세요.",
    "빨리보내주세요.", "빨리보내주세요~",
    "부재시 집앞에 놔둬주세오", "부재시 집앞에 놔둬주세요", "부재시 문앞에 놔주세용",
    "빠른배송", "부재 시 연락 부탁드려요", "문 앞에 놔주세요", "문앞에 놓아주세요",
    "문앞에놓아주세요", "부재시 문앞에 놔주세요",
}


def _t(r, key):
    v = r.get(key)
    return "" if v is None else str(v).strip()


def normalize_zip(z):
    s = str(z or "").strip()
    return str(int(s)) if s.isdigit() else s          # 웹 도구와 같이 앞자리 0을 뺌


def normalize_text(s):
    return re.sub(r"[\s\u3000\u00A0\u200B]+", "", str(s or "").strip())


def normalize_phone(s):
    return re.sub(r"[^0-9]", "", str(s or ""))


def is_jeju(address):
    s = str(address or "")
    return "제주" in s or "서귀포" in s


def is_important_message(msg):
    t = str(msg or "").strip()
    return t != "" and t not in TRIVIAL_DELIVERY_MESSAGES


# ---------------------------------------------------------------- 계산
def build_courier_rows(rows):
    """원본 줄(상품 단위) → 택배 줄(주문 단위). 같은 수령인·주소·전화·주문번호는 한 줄."""
    seen, out = set(), []
    for r in rows:
        recipient = _t(r, "수령인")
        zip_ = normalize_zip(_t(r, "우편번호"))
        raw_addr = _t(r, "주소")
        address = raw_addr.replace("~", "-")          # 롯데택배 시스템이 '~'를 인식 못함
        tel1, tel2 = _t(r, "전화번호"), _t(r, "핸드폰")
        order_no = _t(r, "주문번호")
        note = "0" if str(r.get("비고") or "").strip() == "" else str(r.get("비고"))
        key = "|".join([recipient, zip_, address, tel1, tel2, order_no])
        if key in seen:
            continue
        seen.add(key)
        out.append({"recipient": recipient, "zip": zip_, "address": address, "tel1": tel1, "tel2": tel2,
                    "orderNo": order_no, "note": note, "addressChanged": "~" in raw_addr, "extra": False})
    return out


def clean_zip(z):
    """시트에서 숫자로 바뀌어 온 우편번호(-6017, 6017.0 등)를 주문과 같은 형식으로: 숫자만, 앞자리 0 제외, 글자로"""
    s = str(z if z is not None else "").strip()
    if re.fullmatch(r"-?\d+\.0+", s):
        s = s.split(".")[0]
    digits = re.sub(r"\D", "", s)
    return normalize_zip(digits) if digits else ""


def extra_to_courier_rows(items):
    return [{"recipient": it.get("name") or "", "zip": clean_zip(it.get("zip")),
             "address": str(it.get("address") or "").replace("~", "-"),
             "tel1": it.get("tel1") or "", "tel2": it.get("tel2") or "",
             "orderNo": it.get("memo") or "추가배송", "note": it.get("message") or "0",
             "addressChanged": "~" in str(it.get("address") or ""), "extra": True} for it in items]


def compute_highlights(rows):
    """이름+주소가 같은 사람(공백 무시) → 노란색 + 묶음배송 요약"""
    by_name = {}
    for i, r in enumerate(rows):
        by_name.setdefault(normalize_text(r["recipient"]), []).append(i)
    yellow, summary = set(), []
    for idxs in by_name.values():
        if len(idxs) < 2:
            continue
        by_addr = {}
        for i in idxs:
            by_addr.setdefault(normalize_text(rows[i]["address"]), []).append(i)
        for a in by_addr.values():
            if len(a) >= 2:
                yellow.update(a)
                summary.append({"name": rows[a[0]]["recipient"], "count": len(a),
                                "orderNos": [rows[i]["orderNo"] for i in a]})
    return yellow, summary


def compute_blacklist(rows, items):
    """이름+주소 일치 또는 휴대폰 일치 → {줄번호: {level, reasons}}"""
    by_na, by_ph = {}, {}
    for b in items:
        entry = {"level": "경계대상" if b.get("level") == "경계대상" else "블랙리스트",
                 "reason": str(b.get("reason") or "").strip()}
        n, a, p = normalize_text(b.get("name")), normalize_text(b.get("address")), normalize_phone(b.get("phone"))
        if n and a:
            by_na.setdefault(n + "||" + a, []).append(entry)
        if p:
            by_ph.setdefault(p, []).append(entry)
    out = {}
    for i, r in enumerate(rows):
        hits = list(by_na.get(normalize_text(r["recipient"]) + "||" + normalize_text(r["address"]), []))
        for p in (normalize_phone(r["tel1"]), normalize_phone(r["tel2"])):
            if p and p in by_ph:
                hits += by_ph[p]
        if hits:
            level = "블랙리스트" if any(h["level"] == "블랙리스트" for h in hits) else "경계대상"
            reasons = list(dict.fromkeys(h["reason"] for h in hits if h["reason"]))
            out[i] = {"level": level, "reasons": reasons}
    return out


def analyze(order_rows, extra_items, blacklist_items):
    courier = build_courier_rows(order_rows)
    extra = extra_to_courier_rows(extra_items)
    allrows = courier + extra
    yellow, summary = compute_highlights(allrows)
    bl = compute_blacklist(allrows, blacklist_items)
    return {
        "rows": allrows, "extra": extra, "yellow": yellow, "bundle": summary, "blacklist": bl,
        "jeju": [r for r in allrows if is_jeju(r["address"])],
        "tilde": [r["recipient"] for r in allrows if r["addressChanged"]],
        "messages": [r for r in courier if is_important_message(r["note"])],
    }


def summary_texts(a):
    """화면의 확인 상자에 넣을 글 (웹 도구와 같은 형식)"""
    return {
        "묶음배송 확인": "\n".join(f"{g['name']} / {g['count']}건 묶음 / {' / '.join(g['orderNos'])}" for g in a["bundle"]),
        "제주도 확인": "\n".join(f"{r['orderNo']} / {r['recipient']} / {r['tel1']} / {r['tel2']}" for r in a["jeju"]),
        "주소 '~' 변경 확인": "\n".join(a["tilde"]),
        "중요 배송메세지 확인": "\n".join(f"{r['orderNo']} / {r['note']}" for r in a["messages"]),
        "추가배송 포함": "\n".join(f"{r['recipient']} / {r['address']} / {r['tel1'] or r['tel2'] or ''} / {r['orderNo'] or ''}"
                              for r in a["extra"]),
        "블랙리스트·경계대상": "\n".join(
            f"[{b['level']}] {a['rows'][i]['orderNo']} / {a['rows'][i]['recipient']}"
            + (f" / {' / '.join(b['reasons'])}" if b["reasons"] else "") for i, b in sorted(a["blacklist"].items())),
    }


def popup_text(a):
    """웹 도구의 알림 창과 같은 내용 (묶음배송·제주도·주소변경)"""
    parts = []
    if a["bundle"]:
        parts.append("묶음배송: " + ", ".join(f"{g['name']} 묶음 {g['count']}건" for g in a["bundle"]))
    if a["jeju"]:
        parts.append("제주도: " + ", ".join(r["recipient"] for r in a["jeju"]))
    if a["tilde"]:
        parts.append("주소변경: " + ", ".join(a["tilde"]))
    return "\n".join(parts)


# ---------------------------------------------------------------- 엑셀 쓰기
def write_courier_xlsx(a, path):
    from openpyxl import Workbook
    from openpyxl.comments import Comment
    from openpyxl.styles import Alignment, Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = "택배발송"
    ws.append(COURIER_HEADERS)
    for c in ws[1]:
        c.font = Font(bold=True)
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    ws.row_dimensions[1].height = 32
    for r in a["rows"]:
        ws.append([r["recipient"], r["zip"], r["address"], r["tel1"], r["tel2"], r["orderNo"], 1, None, None,
                   r["note"], "신용", None, SENDER["name"], SENDER["tel1"], SENDER["tel2"], SENDER["addr"]])
    from openpyxl.utils import get_column_letter
    for i, w in enumerate(COURIER_WIDTHS, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    fill = lambda argb: PatternFill("solid", fgColor=argb)
    for i in range(len(a["rows"])):
        row = i + 2
        bl = a["blacklist"].get(i)
        if bl:                                           # 블랙리스트·경계대상이 최우선
            f = fill(RED if bl["level"] == "블랙리스트" else BLUE)
            for col in (1, 3):
                ws.cell(row, col).fill = f
                ws.cell(row, col).font = Font(color="FFFFFFFF", bold=True)
            if bl["reasons"]:
                ws.cell(row, 1).comment = Comment(f"[{bl['level']}] " + " / ".join(bl["reasons"]), "주문서 출력 관리")
        elif i in a["yellow"]:
            ws.cell(row, 1).fill = fill(YELLOW)
            ws.cell(row, 3).fill = fill(YELLOW)
    wb.save(path)


# ---------------------------------------------------------------- 구글 시트 (웹 도구와 같은 Apps Script 주소)
class SheetError(Exception):
    pass


def sheet_call(url, action, retries=3, **params):
    """웹 도구와 같은 방식(JSONP)으로 Apps Script 호출. 일시적인 오류는 자동으로 다시 시도."""
    if not url:
        raise SheetError("구글 시트 연결 주소가 설정되지 않았습니다 (관리 탭).")
    last = ""
    for attempt in range(retries):
        cb = f"cb_{time.time_ns()}"
        try:
            r = requests.get(url, params={**params, "action": action, "callback": cb, "_": time.time_ns()},
                             timeout=30, headers={"Cache-Control": "no-cache"})
        except requests.RequestException as e:
            last = f"인터넷 연결 오류 ({type(e).__name__})"
            time.sleep(1.5 * (attempt + 1))
            continue
        if "accounts.google.com" in r.url or "ServiceLogin" in r.text[:3000]:
            raise SheetError("구글 로그인이 필요한 배포 설정입니다. Apps Script → 배포 관리에서 "
                             "액세스 권한을 '모든 사용자'로 바꿔 주세요.")
        if r.status_code != 200:
            last = f"구글 응답 오류 {r.status_code}"
            time.sleep(1.5 * (attempt + 1))
            continue
        text = r.text.strip()
        m = re.match(r"^[\w.]+\((.*)\);?\s*$", text, re.S)
        try:
            data = json.loads(m.group(1) if m else text)
        except (ValueError, AttributeError):
            last = "응답을 읽을 수 없음"
            time.sleep(1.5 * (attempt + 1))
            continue
        if not data.get("ok", False):
            raise SheetError(f"구글 시트 오류: {data.get('error') or '알 수 없음'}")
        return data
    raise SheetError(f"구글 시트 연결 실패 ({retries}번 시도): {last}. 잠시 후 다시 시도해 주세요.")
